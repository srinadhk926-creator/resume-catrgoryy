"""
ResumeAI - Resume Document Parser
=================================
Safely parses and extracts raw text from PDF and DOCX resume documents
using PyMuPDF (fitz) and python-docx.
"""

import io
import re
import logging
from pathlib import Path
from typing import Tuple, Optional, BinaryIO

import fitz  # PyMuPDF
import docx

logger = logging.getLogger("ResumeAI-ResumeParser")

ALLOWED_EXTENSIONS = {"pdf", "docx"}
DEFAULT_MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


class ResumeParserError(Exception):
    """Base exception for resume parsing errors."""
    pass


class UnsupportedFileTypeError(ResumeParserError):
    """Raised when file extension or magic bytes do not match supported types."""
    pass


class FileTooLargeError(ResumeParserError):
    """Raised when file size exceeds configured maximum."""
    pass


class EmptyFileError(ResumeParserError):
    """Raised when file is 0 bytes or extracted text is blank."""
    pass


class CorruptedFileError(ResumeParserError):
    """Raised when file is malformed or cannot be opened by the parser."""
    pass


def sanitize_filename(filename: str) -> str:
    """Removes path traversals and dangerous characters from filenames."""
    filename = Path(filename).name
    # Keep only alphanumeric, hyphens, underscores, dots, and spaces
    filename = re.sub(r"[^\w\.\-\s]", "_", filename)
    return filename[:200]


def detect_file_extension(filename: str) -> str:
    """Extracts and lowercases the file extension without dot."""
    if "." not in filename:
        return ""
    return filename.rsplit(".", 1)[1].lower()


def validate_file_content(file_bytes: bytes, extension: str) -> None:
    """Performs magic byte inspection to verify real file format."""
    if len(file_bytes) == 0:
        raise EmptyFileError("Uploaded file is empty (0 bytes).")

    if extension == "pdf":
        # PDF files must start with %PDF-
        if not file_bytes.startswith(b"%PDF-"):
            raise UnsupportedFileTypeError(
                "File content is not a valid PDF (missing %PDF- header)."
            )
    elif extension == "docx":
        # DOCX is a zip archive and starts with standard zip magic bytes PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise UnsupportedFileTypeError(
                "File content is not a valid Word DOCX document (invalid zip container)."
            )
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file extension '{extension}'. Only PDF and DOCX files are supported."
        )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts textual content from PDF byte stream using PyMuPDF."""
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as err:
        raise CorruptedFileError(f"Failed to open PDF document: {err}") from err

    extracted_pages = []
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text and text.strip():
                extracted_pages.append(text.strip())
    except Exception as err:
        raise CorruptedFileError(f"Error reading PDF page contents: {err}") from err
    finally:
        doc.close()

    full_text = "\n\n".join(extracted_pages).strip()
    if not full_text:
        raise EmptyFileError(
            "PDF contains no selectable or extractable text (it may be a scanned image or empty)."
        )
    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX byte stream using python-docx."""
    try:
        file_stream = io.BytesIO(file_bytes)
        doc = docx.Document(file_stream)
    except Exception as err:
        raise CorruptedFileError(f"Failed to open DOCX document: {err}") from err

    paragraphs = []
    try:
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                paragraphs.append(p.text.strip())
        # Also extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        paragraphs.append(cell.text.strip())
    except Exception as err:
        raise CorruptedFileError(f"Error reading DOCX paragraphs: {err}") from err

    full_text = "\n\n".join(paragraphs).strip()
    if not full_text:
        raise EmptyFileError("DOCX document contains no readable text.")
    return full_text


def parse_resume_file(
    file_stream_or_bytes: BinaryIO | bytes,
    filename: str,
    max_file_size: int = DEFAULT_MAX_FILE_SIZE,
) -> Tuple[str, str]:
    """
    Main parser entrypoint. Validates size, extension, magic bytes, and extracts text.

    Args:
        file_stream_or_bytes: A file-like object or raw bytes.
        filename: Name of the uploaded file.
        max_file_size: Maximum allowed upload size in bytes.

    Returns:
        Tuple of (clean_extracted_text, safe_filename)

    Raises:
        UnsupportedFileTypeError, FileTooLargeError, EmptyFileError, CorruptedFileError
    """
    safe_name = sanitize_filename(filename)
    extension = detect_file_extension(safe_name)

    if extension not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"File '{safe_name}' has unsupported format. Allowed formats: {', '.join(sorted(ALLOWED_EXTENSIONS)).upper()}."
        )

    # Read bytes safely
    if isinstance(file_stream_or_bytes, bytes):
        file_bytes = file_stream_or_bytes
    else:
        file_bytes = file_stream_or_bytes.read()

    # Enforce file size limit
    if len(file_bytes) > max_file_size:
        max_mb = max_file_size / (1024 * 1024)
        actual_mb = round(len(file_bytes) / (1024 * 1024), 2)
        raise FileTooLargeError(
            f"File '{safe_name}' exceeds maximum allowed size ({actual_mb} MB > {max_mb:.1f} MB)."
        )

    # Validate header magic bytes
    validate_file_content(file_bytes, extension)

    # Extract text based on file format
    if extension == "pdf":
        extracted_text = extract_text_from_pdf(file_bytes)
    elif extension == "docx":
        extracted_text = extract_text_from_docx(file_bytes)
    else:
        raise UnsupportedFileTypeError(f"Unsupported format: {extension}")

    return extracted_text, safe_name


def redact_pii(text: str, max_chars: int = 500) -> str:
    """
    Safely bounds extracted resume text and masks common PII patterns
    including emails, phone numbers, and street addresses.
    """
    if not text:
        return ""
    bounded = text[:max_chars].strip()

    # Mask email addresses
    bounded = re.sub(
        r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+",
        "[REDACTED EMAIL]",
        bounded,
    )
    # Mask standard phone numbers (international and national patterns)
    bounded = re.sub(
        r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}",
        "[REDACTED PHONE]",
        bounded,
    )
    # Mask street addresses
    bounded = re.sub(
        r"\b\d{1,5}\s+[A-Za-z0-9\.\s]+(?:\bStreet|\bSt|\bAvenue|\bAve|\bRoad|\bRd|\bBoulevard|\bBlvd|\bDrive|\bDr|\bLane|\bLn|\bSuite|\bApt)\b",
        "[REDACTED ADDRESS]",
        bounded,
        flags=re.IGNORECASE,
    )

    if len(text) > max_chars:
        bounded += "..."

    return bounded
